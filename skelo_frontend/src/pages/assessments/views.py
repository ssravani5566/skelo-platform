import re
from django.shortcuts import get_object_or_404
from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from accounts.permissions import IsFacultyOrAdmin
from core.api import paginated_response
from students.services import get_student_profile
from students.models import StudentProfile
from .models import Assessment, AssessmentQuestion, Question, StudentAssessment, StudentAnswer
from .serializers import AssessmentDetailSerializer, AssessmentListSerializer, AttemptSerializer, StudentAnswerSerializer
from .serializers_admin import QuestionAdminSerializer, AssessmentQuestionSerializer
from .services import get_or_create_attempt, is_open, submit_attempt

class HealthView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self,request): return Response({"status":"ok"})


def _student_is_allowed(assessment, student):
    assigned = assessment.assigned_students.all()
    return not assigned.exists() or assigned.filter(pk=student.pk).exists()


class AssessmentListView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self,request):
        qs=Assessment.objects.prefetch_related("assessment_questions__question", "assigned_students__user").order_by("-created_at")
        if request.user.role=="STUDENT": qs=qs.filter(is_published=True).filter(Q(assigned_students__isnull=True) | Q(assigned_students__pk=get_student_profile(request.user).pk)).distinct()
        student=get_student_profile(request.user) if request.user.role=="STUDENT" else None
        data=AssessmentListSerializer(qs,many=True,context={"student":student}).data
        status_filter=request.query_params.get("status")
        if status_filter: data=[x for x in data if x["status"].lower()==status_filter.lower()]
        q=request.query_params.get("q","").strip().lower()
        if q: data=[x for x in data if q in x["title"].lower() or q in x["description"].lower()]
        return paginated_response(request,data)
    def post(self,request):
        if request.user.role not in ("ADMIN","FACULTY"): return Response({"detail":"Faculty or admin access required."},status=403)
        from .serializers_admin import AssessmentAdminSerializer
        s=AssessmentAdminSerializer(data=request.data); s.is_valid(raise_exception=True); a=s.save(created_by=request.user); return Response(AssessmentDetailSerializer(a).data,status=201)

class AssessmentDetailView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self,request,pk):
        qs=Assessment.objects.prefetch_related("assessment_questions__question", "assigned_students__user")
        if request.user.role=="STUDENT":
            student = get_student_profile(request.user)
            qs=qs.filter(is_published=True).filter(
                Q(assigned_students__isnull=True) | Q(assigned_students__pk=student.pk)
            ).distinct()
        a=get_object_or_404(qs,pk=pk)
        student=get_student_profile(request.user) if request.user.role=="STUDENT" else None
        return Response(AssessmentDetailSerializer(a,context={"student":student}).data)
    def patch(self,request,pk):
        if request.user.role not in ("ADMIN","FACULTY"): return Response({"detail":"Faculty or admin access required."},status=403)
        from .serializers_admin import AssessmentAdminSerializer
        a=get_object_or_404(Assessment,pk=pk); s=AssessmentAdminSerializer(a,data=request.data,partial=True); s.is_valid(raise_exception=True); s.save(); return Response(AssessmentDetailSerializer(a).data)
    def delete(self,request,pk):
        if request.user.role not in ("ADMIN","FACULTY"): return Response({"detail":"Faculty or admin access required."},status=403)
        get_object_or_404(Assessment,pk=pk).delete(); return Response(status=204)

class AssessmentStartView(APIView):
    permission_classes=[IsAuthenticated]
    def post(self,request,pk):
        if request.user.role!="STUDENT": return Response({"detail":"Student access required."},status=403)
        a=get_object_or_404(Assessment,pk=pk,is_published=True)
        student = get_student_profile(request.user)
        if not _student_is_allowed(a, student):
            return Response({"detail":"This assessment is not assigned to you."},status=403)
        if not is_open(a): return Response({"detail":"Assessment is not open."},status=400)
        attempt=get_or_create_attempt(get_student_profile(request.user),a)
        return Response(AttemptSerializer(attempt).data)

class AttemptView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self,request,pk):
        attempt=get_object_or_404(StudentAssessment,pk=pk,student=get_student_profile(request.user))
        return Response(AttemptSerializer(attempt).data)

    def post(self,request,pk):
        attempt=get_object_or_404(StudentAssessment,pk=pk,student=get_student_profile(request.user))
        if attempt.status=="EVALUATED": return Response({"detail":"Attempt already submitted."},status=400)
        qid=request.data.get("question"); answer=request.data.get("answer_text","")
        aq=get_object_or_404(AssessmentQuestion,assessment=attempt.assessment,question_id=qid)
        obj,_=StudentAnswer.objects.update_or_create(student_assessment=attempt,question=aq.question,defaults={"answer_text":answer})
        return Response(StudentAnswerSerializer(obj).data)

    def put(self,request,pk):
        attempt=get_object_or_404(StudentAssessment,pk=pk,student=get_student_profile(request.user))
        if attempt.status=="EVALUATED": return Response({"detail":"Attempt already submitted."},status=400)
        if isinstance(request.data.get("answers"),list):
            for row in request.data["answers"]:
                aq=get_object_or_404(AssessmentQuestion,assessment=attempt.assessment,question_id=row.get("question"))
                StudentAnswer.objects.update_or_create(student_assessment=attempt,question=aq.question,defaults={"answer_text":row.get("answer_text","")})
        return Response(AttemptSerializer(attempt).data)

    def patch(self,request,pk): return self.post(request,pk)

class AttemptSubmitView(APIView):
    permission_classes=[IsAuthenticated]
    def post(self,request,pk):
        attempt=get_object_or_404(StudentAssessment,pk=pk,student=get_student_profile(request.user))
        if attempt.status=="EVALUATED": return Response(AttemptSerializer(attempt).data)
        return Response(AttemptSerializer(submit_attempt(attempt)).data)

class AttemptResultView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self,request,pk):
        attempt=get_object_or_404(StudentAssessment,pk=pk,student=get_student_profile(request.user),status="EVALUATED")
        return Response(AttemptSerializer(attempt).data)

class AssessmentQuestionAdminView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self,request,assessment_id):
        if request.user.role not in ("ADMIN","FACULTY"): return Response({"detail":"Faculty or admin access required."},status=403)
        a=get_object_or_404(Assessment,pk=assessment_id)
        rows=a.assessment_questions.select_related("question").order_by("order")
        return Response(AssessmentQuestionSerializer(rows,many=True).data)
    def post(self,request,assessment_id):
        if request.user.role not in ("ADMIN","FACULTY"): return Response({"detail":"Faculty or admin access required."},status=403)
        a=get_object_or_404(Assessment,pk=assessment_id)
        payload=request.data.copy()
        question_id=payload.pop("question_id",None)
        if question_id:
            q=get_object_or_404(Question,pk=question_id)
            qs=QuestionAdminSerializer(q,data=payload,partial=True)
        else:
            qs=QuestionAdminSerializer(data=payload)
        qs.is_valid(raise_exception=True); q=qs.save()
        order=int(payload.get("order") or a.assessment_questions.count()+1)
        marks=payload.get("marks")
        obj,_=AssessmentQuestion.objects.update_or_create(assessment=a,question=q,defaults={"order":order,"marks":marks})
        a.total_marks=sum((x.marks if x.marks is not None else x.question.marks) for x in a.assessment_questions.select_related("question"))
        a.save(update_fields=["total_marks"])
        return Response(AssessmentQuestionSerializer(obj).data,status=201)

    def patch(self,request,assessment_id):
        if request.user.role not in ("ADMIN","FACULTY"): return Response({"detail":"Faculty or admin access required."},status=403)
        aq=get_object_or_404(AssessmentQuestion,pk=request.data.get("assessment_question_id"),assessment_id=assessment_id)
        payload=request.data.copy(); payload.pop("assessment_question_id",None)
        if "order" in payload or "marks" in payload:
            if "order" in payload: aq.order=int(payload["order"])
            if "marks" in payload: aq.marks=payload["marks"]
            aq.save(update_fields=["order","marks"])
        qdata={k:v for k,v in payload.items() if k!="order" and k!="marks"}
        if qdata:
            qs=QuestionAdminSerializer(aq.question,data=qdata,partial=True); qs.is_valid(raise_exception=True); qs.save()
        aq.refresh_from_db()
        aq.assessment.total_marks=sum((x.marks if x.marks is not None else x.question.marks) for x in aq.assessment.assessment_questions.select_related("question"))
        aq.assessment.save(update_fields=["total_marks"])
        return Response(AssessmentQuestionSerializer(aq).data)

    def delete(self,request,assessment_id):
        if request.user.role not in ("ADMIN","FACULTY"): return Response({"detail":"Faculty or admin access required."},status=403)
        aq=get_object_or_404(AssessmentQuestion,pk=request.data.get("assessment_question_id"),assessment_id=assessment_id)
        q=aq.question; assessment=aq.assessment; aq.delete()
        if not AssessmentQuestion.objects.filter(question=q).exists(): q.delete()
        assessment.total_marks=sum((x.marks if x.marks is not None else x.question.marks) for x in assessment.assessment_questions.select_related("question"))
        assessment.save(update_fields=["total_marks"])
        return Response(status=204)




class AssessmentStudentsView(APIView):
    permission_classes=[IsAuthenticated]

    def get(self, request):
        if request.user.role not in ("ADMIN","FACULTY"):
            return Response({"detail":"Faculty or admin access required."},status=403)

        students = StudentProfile.objects.select_related("user").filter(
            user__role="STUDENT",
            user__is_active=True,
        ).order_by("user__first_name","user__last_name","user__email")

        return Response([
            {
                "id": student.id,
                "user_id": student.user_id,
                "name": student.user.get_full_name() or student.user.email,
                "email": student.user.email,
                "register_number": student.register_number,
                "department": student.department,
                "year": student.year,
            }
            for student in students
        ])


class AssessmentQuestionImportView(APIView):
    permission_classes=[IsAuthenticated]

    def post(self, request, assessment_id):
        if request.user.role not in ("ADMIN","FACULTY"):
            return Response({"detail":"Faculty or admin access required."},status=403)

        assessment = get_object_or_404(Assessment, pk=assessment_id)
        uploaded = request.FILES.get("file")

        if not uploaded:
            return Response({"detail":"Please upload an Excel .xlsx file."},status=400)

        if not uploaded.name.lower().endswith(".xlsx"):
            return Response({"detail":"Only .xlsx Excel files are supported."},status=400)

        try:
            from openpyxl import load_workbook
        except ImportError:
            return Response({
                "detail":"openpyxl is not installed. Run: pip install openpyxl"
            },status=500)

        def clean(value):
            return "" if value is None else str(value).strip()

        def header_key(value):
            return re.sub(r"[^a-z0-9]+", "_", clean(value).lower()).strip("_")

        def answer_value(value):
            value = clean(value)
            upper = value.upper()
            if upper in {"A","B","C","D"}:
                return upper
            return value

        try:
            workbook = load_workbook(uploaded, read_only=True, data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))

            if not rows:
                return Response({"detail":"The Excel file is empty."},status=400)

            headers = [header_key(v) for v in rows[0]]
            aliases = {
                "question": {"questions","question","question_text","text"},
                "option_a": {"option_a","a","option1","option_1"},
                "option_b": {"option_b","b","option2","option_2"},
                "option_c": {"option_c","c","option3","option_3"},
                "option_d": {"option_d","d","option4","option_4"},
                "answer": {"answer","answers","correct_answer","correct"},
                "marks": {"mark","marks","mark_","marks_"},
                "explanation": {"explain_answer","explanation","explain","solution"},
                "input_format": {"input_format","input"},
                "output_format": {"output_format","output"},
                "difficulty": {"difficulty_level","difficulty","level"},
            }

            indexes = {}
            for target, names in aliases.items():
                for i, header in enumerate(headers):
                    if header in names:
                        indexes[target] = i
                        break

            if "question" not in indexes:
                return Response({
                    "detail":"Question column was not found. Use 'Questions**' or 'Question'."
                },status=400)

            imported = 0
            errors = []

            with transaction.atomic():
                for row_number, row in enumerate(rows[1:], start=2):
                    get_value = lambda key: (
                        clean(row[indexes[key]])
                        if key in indexes and indexes[key] < len(row)
                        else ""
                    )

                    text = get_value("question")
                    if not text:
                        continue

                    option_a = get_value("option_a")
                    option_b = get_value("option_b")
                    option_c = get_value("option_c")
                    option_d = get_value("option_d")
                    answer = answer_value(get_value("answer"))

                    if option_a or option_b or option_c or option_d:
                        qtype = "MCQ"
                        if answer.upper() not in {"A","B","C","D"}:
                            errors.append(
                                f"Row {row_number}: MCQ answer must be A/B/C/D."
                            )
                            continue
                    else:
                        # Your QuestionPaper_1999 format has no A/B/C/D columns.
                        # When the explanation contains exported source code and
                        # an Input Format column exists, treat it as a coding
                        # question instead of losing it as a short-answer item.
                        explanation_value = get_value("explanation")
                        if (
                            "code:" in explanation_value.lower()
                            and get_value("input_format")
                        ):
                            qtype = "CODING"
                        else:
                            qtype = "SHORT"

                    try:
                        marks = int(float(get_value("marks") or 1))
                    except ValueError:
                        marks = 1

                    question = Question.objects.create(
                        text=text,
                        question_type=qtype,
                        option_a=option_a,
                        option_b=option_b,
                        option_c=option_c,
                        option_d=option_d,
                        correct_answer=answer,
                        explanation=get_value("explanation"),
                        difficulty=get_value("difficulty"),
                        marks=max(marks,1),
                        input_format=get_value("input_format"),
                        output_format=get_value("output_format"),
                        test_cases=(
                            [{"input": "", "output": answer}]
                            if qtype == "CODING" and answer
                            else []
                        ),
                    )

                    AssessmentQuestion.objects.create(
                        assessment=assessment,
                        question=question,
                        order=assessment.assessment_questions.count()+1,
                        marks=max(marks,1),
                    )

                    imported += 1

                if errors:
                    raise ValueError(" ".join(errors))

                assessment.total_marks = sum(
                    (aq.marks if aq.marks is not None else aq.question.marks)
                    for aq in assessment.assessment_questions.select_related("question")
                )
                assessment.save(update_fields=["total_marks"])

            return Response({
                "message": f"{imported} questions imported successfully.",
                "imported": imported,
                "total_marks": assessment.total_marks,
            }, status=201)

        except ValueError as exc:
            return Response({"detail":str(exc)},status=400)
        except Exception as exc:
            return Response({"detail":f"Unable to read Excel file: {exc}"},status=400)


class AdminAssessmentResultsView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self, request):
        if request.user.role != "ADMIN":
            return Response({"detail":"Admin access required."}, status=403)
        rows = StudentAssessment.objects.select_related("student__user", "assessment").prefetch_related("answers").order_by("-submitted_at", "-id")
        result=[]
        for attempt in rows:
            answers=list(attempt.answers.all())
            correct=sum(1 for a in answers if a.is_correct is True)
            wrong=sum(1 for a in answers if a.is_correct is False)
            total=attempt.assessment.assessment_questions.count()
            result.append({
                "id":attempt.id,
                "student_id":attempt.student.id,
                "student_name":attempt.student.user.get_full_name() or attempt.student.user.username,
                "student_email":attempt.student.user.email,
                "assessment_id":attempt.assessment.id,
                "assessment_title":attempt.assessment.title,
                "assessment_type":attempt.assessment.assessment_type,
                "status":attempt.status,
                "total_questions":total,
                "answered":len(answers),
                "correct_answers":correct,
                "wrong_answers":wrong,
                "unanswered":max(total-len(answers),0),
                "score":float(attempt.score),
                "percentage":float(attempt.percentage),
                "submitted_at":attempt.submitted_at,
            })
        return Response(result)


class AssessmentCodingAnswerView(APIView):
    permission_classes=[IsAuthenticated]
    def post(self,request,pk):
        if request.user.role!="STUDENT": return Response({"detail":"Student access required."},status=403)
        attempt=get_object_or_404(StudentAssessment,pk=pk,student=get_student_profile(request.user))
        if attempt.status=="EVALUATED": return Response({"detail":"Attempt already submitted."},status=400)
        q=get_object_or_404(Question,pk=request.data.get("question"),question_type="CODING",assessment_questions__assessment=attempt.assessment)
        language=str(request.data.get("language") or q.coding_language or "PYTHON").upper()
        source=str(request.data.get("source_code") or "")
        from coding.services import run_code
        from types import SimpleNamespace
        cases=[SimpleNamespace(input_data=str(x.get("input", "")),expected_output=str(x.get("output", ""))) for x in (q.test_cases or [])]
        status_code,error,results=run_code(language,source,cases)
        payload={"language":language,"source_code":source,"status":status_code,"error":error,"passed":sum(1 for r in results if r["passed"]),"total":len(cases),"tests":[{"passed":r["passed"],"actual_output":r.get("actual_output",""),"runtime_ms":r.get("runtime_ms",0),"error":r.get("error_message","")} for r in results]}
        answer,_=StudentAnswer.objects.update_or_create(student_assessment=attempt,question=q,defaults={"answer_text":__import__("json").dumps(payload),"is_correct":status_code=="ACCEPTED"})
        return Response({"status":status_code,"error":error,"passed":payload["passed"],"total":payload["total"],"tests":payload["tests"],"answer":StudentAnswerSerializer(answer).data})

